from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_PATH = Path(__file__).resolve().parents[1] / 'Book1.xlsx'
OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'supabase'
JSON_OUTPUT = OUTPUT_DIR / 'master_data_import.json'
SQL_OUTPUT = OUTPUT_DIR / 'master_data_import.sql'


def normalize_project_name(name: str) -> str:
    value = (name or '').strip()
    if value.endswith(' '):
        value = value.rstrip()
    return value


def parse_workbook():
    wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    records = []

    for sheet in wb.worksheets:
        project_name = normalize_project_name(sheet.title)
        if not project_name:
            continue

        current_parameter = None
        for row in sheet.iter_rows(values_only=True):
            if len(row) < 2:
                continue

            col_a = (row[0] or '').strip() if row[0] is not None else ''
            col_b = (row[1] or '').strip() if row[1] is not None else ''

            if not col_a and not col_b:
                continue

            if col_a and col_a.lower() not in {'category', 'product variant name', 'warehouse inventory tracking'}:
                current_parameter = col_a
                continue

            if current_parameter and col_b:
                records.append({
                    'project': project_name,
                    'parameter': current_parameter,
                    'item': col_b,
                })

    return records


def build_sql(records):
    project_names = sorted({r['project'] for r in records})
    sql_lines = ['-- Auto-generated import for project master data extracted from Book1.xlsx', '']

    for project_name in project_names:
        code = project_name
        sql_lines.append(
            f"INSERT INTO public.projects (name, code, description, status) VALUES ('{project_name}', '{code}', 'Imported from Book1.xlsx', 'active') ON CONFLICT (name) DO NOTHING;"
        )

    sql_lines.append('')

    parameters_by_project = {}
    for record in records:
        parameters_by_project.setdefault(record['project'], set()).add(record['parameter'])

    for project_name, parameters in sorted(parameters_by_project.items()):
        for parameter_name in sorted(parameters):
            sql_lines.append(
                f"INSERT INTO public.parameters (project_id, name, description, status) "
                f"SELECT id, '{parameter_name}', 'Imported from Book1.xlsx', 'active' "
                f"FROM public.projects WHERE name = '{project_name}' ON CONFLICT (project_id, name) DO NOTHING;"
            )

    sql_lines.append('')

    for project_name in project_names:
        parameters = [r['parameter'] for r in records if r['project'] == project_name]
        seen_items = set()
        for record in records:
            if record['project'] != project_name:
                continue
            item_name = record['item']
            if item_name in seen_items:
                continue
            seen_items.add(item_name)
            sql_lines.append(
                f"INSERT INTO public.items (project_id, parameter_id, name, description, status) "
                f"SELECT p.id, param.id, '{item_name}', 'Imported from Book1.xlsx', 'active' "
                f"FROM public.projects p JOIN public.parameters param ON param.project_id = p.id "
                f"WHERE p.name = '{project_name}' AND param.name = '{record['parameter']}' ON CONFLICT (project_id, parameter_id, name) DO NOTHING;"
            )

    return '\n'.join(sql_lines) + '\n'


def main():
    records = parse_workbook()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    JSON_OUTPUT.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    SQL_OUTPUT.write_text(build_sql(records), encoding='utf-8')

    print(f'Parsed {len(records)} item records from {WORKBOOK_PATH.name}')
    print(f'JSON written to {JSON_OUTPUT}')
    print(f'SQL written to {SQL_OUTPUT}')
    print('Projects:', sorted({r['project'] for r in records}))


if __name__ == '__main__':
    main()
